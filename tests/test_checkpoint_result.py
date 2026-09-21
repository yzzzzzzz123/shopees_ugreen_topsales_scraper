from copy import deepcopy
from contextlib import nullcontext
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/shopee-ugreen-topsales/scripts'))
import checkpoint_result
from checkpoint_result import build_snapshot, CheckpointSaver
from publish_result import _convert


def payload():
    url = 'https://shopee.ph/product/64922227/123'
    return {'schema_version':'1.0', 'store_url':'https://shopee.ph/ugreen.ph?page=0&sortBy=sales&tab=0',
            'started_at':'2026-09-21T10:00:00+08:00', 'captured_at':'2026-09-21T11:00:00+08:00',
            'pages':[{'page_index':0,'total_pages':1,'url':'https://shopee.ph/ugreen.ph?page=0&sortBy=sales&tab=0',
                      'next_enabled':False,'cards':[{'product_url':url,'title':'列表标题','price_php':0,'monthly_sales_text':'0 Sold/Month'}]}],
            'details':[]}


def detail():
    return {'product_url':'https://shopee.ph/product/64922227/123','title':'=商品字面文本',
            'skus':[{'model_id':'001','sku_props':{'颜色':'黑色'},'sku_image_url':None}],
            'main_image_url':'https://down-ph.img.susercontent.com/file/testimage',
            'secondary_image_urls':[],'models_complete':True,'gallery_complete':True}


class SnapshotTests(unittest.TestCase):
    def test_pending_details_remain_blank_not_zero_and_list_data_are_retained(self):
        result = build_snapshot(payload())
        row = result['sheets'][0][0]
        self.assertEqual(row[3:6], ['64922227','123','列表标题'])
        self.assertEqual(row[6:10], [0,'0 Sold/Month',0,'已展示'])
        self.assertEqual(row[10:13], [None,None,None])
        self.assertIn('未采集', row[14])
        self.assertEqual(result['sheets'][1:3], [[],[]])
        self.assertEqual(result['sheets'][3][-1][1], '未完成')

    def test_complete_detail_uses_real_title_sku_and_gallery(self):
        p = payload(); p['details'] = [detail()]
        result = build_snapshot(p)
        self.assertEqual(result['counts']['detail_count'],1)
        self.assertEqual(result['sheets'][0][0][5], '=商品字面文本')
        self.assertEqual(result['sheets'][1][0][5], '001')
        self.assertEqual(result['sheets'][2][0][5:7], ['主图',1])
        self.assertEqual(result['sheets'][3][-1][1], '未完成')

    def test_duplicate_and_out_of_scope_details_rejected(self):
        for records in ([detail(), detail()], [dict(detail(), product_url='https://shopee.ph/product/64922227/124')]):
            p = payload(); p['details'] = records
            with self.assertRaises(ValueError): build_snapshot(p)

    def test_unverified_models_or_gallery_are_not_written_as_complete(self):
        for field in ('models_complete','gallery_complete'):
            p = payload(); p['details'] = [dict(detail(), **{field:False})]
            with self.assertRaises(ValueError): build_snapshot(p)

    def test_nonfinite_or_boolean_prices_rejected(self):
        for value in (float('nan'), float('inf'), True, -1):
            p = payload(); p['pages'][0]['cards'][0]['price_php'] = value
            with self.assertRaises(ValueError): build_snapshot(p)

    def test_full_publication_still_rejects_missing_details(self):
        with self.assertRaisesRegex(ValueError,'详情缺失'): _convert(payload())

    def test_saved_partial_list_is_a_valid_prefix_not_fake_last_page(self):
        p = payload(); page=p['pages'][0]; page['total_pages']=2; page['next_enabled']=True
        page['cards']=[dict(page['cards'][0], product_url=f'https://shopee.ph/product/64922227/{n}') for n in range(1,31)]
        result = build_snapshot(p)
        self.assertEqual(result['counts']['total_pages'],2)
        self.assertEqual(result['counts']['list_page_count'],1)
        page['next_enabled']=False
        with self.assertRaises(ValueError): build_snapshot(p)


class SaverTests(unittest.TestCase):
    def test_real_renderer_writes_only_the_requested_xlsx(self):
        snapshot = build_snapshot(payload())
        renderer = Path(__file__).resolve().parents[1] / 'skills/shopee-ugreen-topsales/scripts/checkpoint_excel.mjs'
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            output = Path(directory) / 'progress.xlsx'
            proc = subprocess.run(
                ['node', str(renderer), str(output)],
                input=json.dumps(snapshot, ensure_ascii=False, allow_nan=False),
                text=True,
                capture_output=True,
                timeout=120,
            )
            self.assertEqual(proc.returncode, 0, proc.stderr[-1000:])
            self.assertTrue(output.is_file())
            self.assertEqual(sorted(path.name for path in output.parent.iterdir()), [output.name])

    def test_checkpoint_rejects_unknown_result_files_before_rendering(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            root = Path(directory) / 'project'
            daily = root / 'worktrees/20260921_ugreen_topsales'
            result = daily / 'result'
            result.mkdir(parents=True)
            unknown = result / 'notes.txt'
            unknown.write_text('preserve', encoding='utf-8')
            final = result / 'ugreen_topsales.xlsx'
            with patch.object(checkpoint_result, '_target_workbook', return_value=final), \
                    patch.object(checkpoint_result, '_exclusive_publish', return_value=nullcontext()):
                with self.assertRaisesRegex(ValueError, '未知文件'):
                    checkpoint_result.write_snapshot(payload(), root, daily)
            self.assertEqual(unknown.read_text(encoding='utf-8'), 'preserve')

    def test_checkpoint_removes_renderer_inspection_sidecar(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            root = Path(directory) / 'project'
            daily = root / 'worktrees/20260921_ugreen_topsales'
            result = daily / 'result'
            result.mkdir(parents=True)
            final = result / 'ugreen_topsales.xlsx'
            sidecars = []

            def render(args, **kwargs):
                candidate = Path(args[-1])
                candidate.write_bytes(b'candidate workbook')
                sidecar = candidate.with_name(candidate.name + '.inspect.ndjson')
                sidecar.write_text('captured rows must not persist', encoding='utf-8')
                sidecars.append(sidecar)
                return subprocess.CompletedProcess(args, 0, stdout='', stderr='')

            with patch.object(checkpoint_result, '_target_workbook', return_value=final), \
                    patch.object(checkpoint_result, '_exclusive_publish', return_value=nullcontext()), \
                    patch.object(checkpoint_result.subprocess, 'run', side_effect=render), \
                    patch.object(checkpoint_result, 'add_hyperlinks'), \
                    patch.object(checkpoint_result, 'validate_snapshot'):
                receipt = checkpoint_result.write_snapshot(payload(), root, daily)
            self.assertEqual(receipt['path'], str(result / 'ugreen_topsales_progress.xlsx'))
            self.assertTrue((result / 'ugreen_topsales_progress.xlsx').is_file())
            self.assertTrue(sidecars)
            self.assertFalse(any(sidecar.exists() for sidecar in sidecars))

    def test_existing_completed_detail_cannot_regress_sku_or_gallery(self):
        p = payload()
        p['details'] = [detail()]
        previous = build_snapshot(p)
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            target = Path(directory) / 'ugreen_topsales_progress.xlsx'
            wb = Workbook()
            wb.remove(wb.active)
            for name, rows in zip(('商品汇总', 'SKU明细', '图片明细', '抓取核验'), previous['sheets']):
                ws = wb.create_sheet(name)
                if name == '抓取核验':
                    ws['A2'] = 'UGREEN Top Sales 抓取核验（采集进度）'
                ws['A4'] = '启动时间：' + previous['started_at']
                for row_index, row in enumerate(rows, 7):
                    for column_index, value in enumerate(row, 1):
                        ws.cell(row_index, column_index, value)
            wb.save(target)
            degraded = deepcopy(previous)
            degraded['sheets'][1] = []
            with self.assertRaisesRegex(ValueError, 'SKU 或图库'):
                checkpoint_result._check_existing(target, degraded)

    def test_boundaries_and_tail_are_cumulative(self):
        current={'details':[]}; counts=[]
        def save(p):
            counts.append(len(p['details']))
            return {'detail_count':len(p['details'])}
        saver=CheckpointSaver(lambda:current, save)
        for count in (1,49,50,51,99,100,101):
            current['details']=[None]*count
            saver.maybe_save()
        saver.maybe_save(force=True)
        self.assertEqual(counts,[50,100,101])

    def test_resumed_58_saves_at_100_not_108(self):
        current={'details':[None]*99}; save=Mock(return_value={'detail_count':100})
        saver=CheckpointSaver(lambda:current, save, saved_count=58)
        saver.maybe_save(); save.assert_not_called()
        current['details'].append(None); saver.maybe_save(); save.assert_called_once()

    def test_failed_save_does_not_advance_counter(self):
        saver=CheckpointSaver(lambda:{'details':[None]*50}, Mock(side_effect=OSError('full')))
        with self.assertRaises(OSError): saver.maybe_save()
        self.assertEqual(saver.saved_count,0)

    def test_lost_memory_is_not_allowed_to_overwrite_progress(self):
        save=Mock(); saver=CheckpointSaver(lambda:{'details':[]},save,saved_count=50)
        with self.assertRaises(ValueError): saver.maybe_save(force=True)
        save.assert_not_called()


if __name__ == '__main__': unittest.main()
