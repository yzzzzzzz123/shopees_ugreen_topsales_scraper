from copy import deepcopy
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/shopee-ugreen-topsales/scripts'))
from checkpoint_result import build_snapshot, HEADERS, validate_snapshot, CheckpointSaver
from checkpoint_resume import read_checkpoint_file, load_checkpoint, pending_cards
from checkpoint_state import embed_state, read_embedded, PART
from test_checkpoint_result import payload, detail


def fixture(path, data):
    snapshot = build_snapshot(data)
    wb = Workbook(); wb.remove(wb.active)
    for index, (name, headers, rows) in enumerate(zip(
            ('商品汇总', 'SKU明细', '图片明细', '抓取核验'), HEADERS, snapshot['sheets'])):
        ws = wb.create_sheet(name)
        ws['A2'] = f'UGREEN Top Sales {name}（采集进度）'
        ws['A3'] = snapshot['status']
        ws['A4'] = '启动时间：' + snapshot['started_at']
        ws['A5'] = '保存时间：' + snapshot['saved_at']
        for col, value in enumerate(headers, 1): ws.cell(6, col, value)
        if not rows: ws.cell(7, 1).number_format = '@'
        for r, row in enumerate(rows, 7):
            for c, value in enumerate(row, 1):
                cell = ws.cell(r, c, value)
                if isinstance(value, str): cell.data_type = 's'
                if isinstance(value, str) and value.startswith('https://'): cell.hyperlink = value
                if (index == 0 and c in (4, 5)) or (index in (1, 2) and c in (3, 4)) or (index == 1 and c == 6):
                    cell.number_format = '@'
    wb.save(path)
    validate_snapshot(path, snapshot)


class RecoveryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1])
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'progress.xlsx'
        self.data = payload()
        self.data['details'] = [detail()]
        card = dict(self.data['pages'][0]['cards'][0], product_url='https://shopee.ph/product/64922227/124')
        self.data['pages'][0]['cards'].append(card)

    def save(self, *, embedded=True):
        fixture(self.path, self.data)
        if embedded: embed_state(self.path, self.data)

    def test_native_roundtrip_retains_original_title_and_delimiter_values(self):
        self.data['details'][0]['skus'][0]['sku_props'] = {'name: with colon': 'value, with comma: exact'}
        self.save()
        recovered = read_checkpoint_file(self.path)
        self.assertEqual(recovered['payload'], self.data)
        self.assertEqual(recovered['saved_count'], 1)
        self.assertEqual(recovered['next_checkpoint_at'], 50)
        self.assertEqual([c['product_url'] for c in recovered['pending_cards']], ['https://shopee.ph/product/64922227/124'])

    def test_custom_part_does_not_change_any_excel_part(self):
        self.save(embedded=False)
        with zipfile.ZipFile(self.path) as z: before = {n: z.read(n) for n in z.namelist() if n.startswith('xl/')}
        embed_state(self.path, self.data)
        with zipfile.ZipFile(self.path) as z: after = {n: z.read(n) for n in z.namelist() if n.startswith('xl/')}
        self.assertEqual(before, after)

    def test_missing_state_requires_explicit_bound_hash(self):
        self.save(embedded=False)
        with self.assertRaisesRegex(ValueError, 'SHA-256'): read_checkpoint_file(self.path)
        with self.assertRaisesRegex(ValueError, '不匹配'): read_checkpoint_file(self.path, legacy_sha256='0' * 64)

    def test_legacy_business_cells_roundtrip_and_source_unchanged(self):
        self.data['details'][0]['skus'][0]['sku_props'] = {'Quantity': 'Honkai: Star Rail*1 pcs'}
        self.save(embedded=False)
        before = self.path.read_bytes()
        state = read_checkpoint_file(self.path, legacy_sha256=hashlib.sha256(before).hexdigest())
        self.assertEqual(build_snapshot(state['payload'])['sheets'], build_snapshot(self.data)['sheets'])
        self.assertEqual(state['origin']['kind'], 'legacy_visible_rows')
        self.assertTrue(state['origin']['warnings'])
        self.assertEqual(self.path.read_bytes(), before)

    def test_legacy_unsplittable_specs_rejected(self):
        self.data['details'][0]['skus'][0]['sku_props'] = {'Color': 'Black, Grey'}
        self.save(embedded=False)
        with self.assertRaisesRegex(ValueError, '不能安全拆解'):
            read_checkpoint_file(self.path, legacy_sha256=hashlib.sha256(self.path.read_bytes()).hexdigest())

    def test_native_preserves_cross_page_duplicates_and_pending_is_identity_based(self):
        data = payload(); first = data['pages'][0]
        first.update(total_pages=2, next_enabled=True)
        base = first['cards'][0]
        first['cards'] = [dict(base, product_url=f'https://shopee.ph/product/64922227/{n}') for n in range(123, 153)]
        data['pages'].append(dict(page_index=1, total_pages=2, next_enabled=False,
            url=data['store_url'].replace('page=0', 'page=1'), cards=[dict(base), dict(base, product_url='https://shopee.ph/product/64922227/200')]))
        data['details'] = [dict(detail(), product_url=base['product_url'] + '?display_model_id=1')]
        self.data = data; self.save()
        result = read_checkpoint_file(self.path)
        self.assertEqual(result['payload'], data)
        self.assertEqual(len(result['pending_cards']), 30)
        self.assertEqual(len(pending_cards(data)), 30)

    def test_legacy_cannot_reconstruct_duplicate_page_occurrences(self):
        self.test_native_preserves_cross_page_duplicates_and_pending_is_identity_based()
        fixture(self.path, self.data)
        with self.assertRaisesRegex(ValueError, '跨页重复'):
            read_checkpoint_file(self.path, legacy_sha256=hashlib.sha256(self.path.read_bytes()).hexdigest())

    def test_state_and_visible_cells_must_agree(self):
        self.save()
        changed = deepcopy(self.data); changed['pages'][0]['cards'][0]['price_php'] = 99
        embed_state(self.path, changed)
        with self.assertRaisesRegex(ValueError, '逐单元格'): read_checkpoint_file(self.path)

    def test_corrupt_state_never_falls_back_to_legacy(self):
        self.save()
        with zipfile.ZipFile(self.path) as z: members = [(i, z.read(i.filename)) for i in z.infolist()]
        with zipfile.ZipFile(self.path, 'w') as z:
            for info, data in members:
                z.writestr(info, data.replace(b'native_capture', b'broken_capture') if info.filename == PART else data)
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        with self.assertRaisesRegex(ValueError, '校验值'):
            read_checkpoint_file(self.path, legacy_sha256=digest)

    def test_resume_run_date_must_match_registered_directory(self):
        self.path = self.path.with_name('ugreen_topsales_progress.xlsx'); self.save()
        with patch('checkpoint_resume._target_workbook', return_value=self.path.with_name('ugreen_topsales.xlsx')):
            with self.assertRaisesRegex(ValueError, '启动日'):
                load_checkpoint(Path(self.directory.name), Path(self.directory.name) / '20260922_ugreen_topsales')

    def test_restored_counter_saves_at_next_absolute_multiple_of_fifty(self):
        current = {'details': [None] * 699}
        calls = []
        saver = CheckpointSaver(lambda: current, lambda p: calls.append(len(p['details'])) or {'detail_count': len(p['details'])}, saved_count=680)
        self.assertIsNone(saver.maybe_save())
        current['details'].append(None); saver.maybe_save()
        self.assertEqual(calls, [700])


if __name__ == '__main__': unittest.main()
