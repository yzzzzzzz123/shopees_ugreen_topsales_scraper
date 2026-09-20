"""发布桥接离线测试：只在系统临时目录创建虚构数据与 Git 仓库，不访问浏览器。"""

from __future__ import annotations

from copy import deepcopy
from contextlib import redirect_stderr, redirect_stdout
import fcntl
import hashlib
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "skills/shopee-ugreen-topsales/scripts"
sys.path.insert(0, str(SCRIPTS))
import excel  # noqa: E402
import publish_result as publisher  # noqa: E402


def image_url(number: int) -> str:
    return f"https://down-ph.img.susercontent.com/file/ph-11134207-offline-{number}"


def fixture(count: int = 2) -> dict:
    cards, details = [], []
    for number in range(1, count + 1):
        url = f"https://shopee.ph/product/64922227/{100 + number}"
        cards.append({"product_url": url, "title": f"离线列表夹具 {number}",
                      "price_php": 10.5, "monthly_sales_text": "1.2K+ Sold/Month" if number == 1 else None})
        details.append({"product_url": url, "title": f"离线详情夹具 {number}",
                        "skus": [{"model_id": str(1000 + number), "sku_props": {"颜色": "离线黑色"},
                                  "sku_image_url": None}], "main_image_url": image_url(number * 2),
                        "secondary_image_urls": [image_url(number * 2 + 1)],
                        "models_complete": True, "gallery_complete": True})
    pages = []
    for index in range((count + 29) // 30):
        total = (count + 29) // 30
        pages.append({"page_index": index, "total_pages": total,
                      "url": f"https://shopee.ph/ugreen.ph?page={index}&sortBy=sales&tab=0",
                      "next_enabled": index + 1 < total, "cards": cards[index * 30:(index + 1) * 30]})
    return {"schema_version": "1.0", "store_url": publisher.STORE_URL,
            "started_at": "2000-01-02T23:59:00+08:00", "captured_at": "2000-01-03T00:01:00+08:00",
            "pages": pages, "details": details}


def template(path: Path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, headers in zip(excel.SHEET_NAMES, (excel.SUMMARY_HEADERS, excel.SKU_HEADERS, excel.IMAGE_HEADERS, excel.AUDIT_HEADERS)):
        sheet = workbook.create_sheet(name)
        for column, text in enumerate(headers, 1):
            sheet.cell(excel.HEADER_ROW, column, text)
    workbook.save(path)
    workbook.close()


class PayloadTests(unittest.TestCase):
    def test_conversion_recomputes_rank_status_and_counts_without_sku_image_fallback(self):
        value = fixture()
        value["details"].reverse()
        result, started = publisher._convert(value)
        first, second = result["products"]
        self.assertEqual(first["global_rank"], 1)
        self.assertEqual(first["source_page"], 0)
        self.assertEqual(first["source_position"], 1)
        self.assertEqual(first["monthly_sales_count_lower_bound"], 1200)
        self.assertEqual(second["monthly_sales_observation"], "not_displayed")
        self.assertIsNone(second["monthly_sales_count_lower_bound"])
        self.assertIsNone(first["skus"][0]["sku_image_url"])
        self.assertEqual(first["title"], "离线详情夹具 1")
        self.assertEqual(result["captured_at"], "2000/01/03 00:01:00")
        self.assertEqual(started.astimezone(publisher.SHANGHAI).day, 2)
        self.assertEqual(result["audit"]["products_without_monthly_sales_display"], 1)

    def test_cross_page_dedup_keeps_first_rank_page_position(self):
        value = fixture(31)
        value["pages"][1]["cards"].insert(0, deepcopy(value["pages"][0]["cards"][0]))
        result, _ = publisher._convert(value)
        self.assertEqual(len(result["products"]), 31)
        self.assertEqual(result["products"][0]["source_page"], 0)
        self.assertEqual(result["products"][-1]["source_position"], 2)
        self.assertEqual(result["audit"]["list_input_product_occurrences"], 32)
        self.assertEqual(result["audit"]["list_duplicate_occurrence_count"], 1)

    def test_monthly_exact_lower_bound_and_not_cumulative(self):
        for text, lower in (("0 Sold/Month", 0), ("1,234 Sold/Month", 1234),
                            ("2M+ Sold/Month", 2000000), ("99+ Sold/Month", 99)):
            self.assertEqual(publisher._monthly(text)[1], lower)
        for text in ("1.2K Sold", "1 Sold/Month junk", "1 Sold/Month\n", "1,23 Sold/Month", "1.5 Sold/Month", "1.0001K Sold/Month", "NaN Sold/Month", ""):
            with self.subTest(text=text), self.assertRaises(ValueError):
                publisher._monthly(text)

    def test_invalid_payload_variants_are_rejected(self):
        def card(value):
            return value["pages"][0]["cards"][0]

        def detail(value):
            return value["details"][0]

        cases = {
            "schema": lambda v: v.update(schema_version="2"),
            "untrusted audit": lambda v: v.update(audit={"detail_failure_count": 0}),
            "naive started": lambda v: v.update(started_at="2000-01-02T12:00:00"),
            "naive captured": lambda v: v.update(captured_at="2000-01-03"),
            "time reversal": lambda v: v.update(captured_at="1999-01-01T00:00:00Z"),
            "empty pages": lambda v: v.update(pages=[]),
            "wrong store": lambda v: v.update(store_url=v["store_url"].replace("ugreen.ph", "other")),
            "missing detail": lambda v: v["details"].pop(),
            "duplicate detail": lambda v: v["details"].append(deepcopy(detail(v))),
            "extra detail": lambda v: detail(v).update(product_url="https://shopee.ph/product/64922227/999999"),
            "page order": lambda v: v["pages"][0].update(page_index=1),
            "page count": lambda v: v["pages"][0].update(total_pages=2),
            "page URL mismatch": lambda v: v["pages"][0].update(url=v["pages"][0]["url"].replace("page=0", "page=1")),
            "wrong sort": lambda v: v["pages"][0].update(url=v["pages"][0]["url"].replace("sales", "price")),
            "next state": lambda v: v["pages"][0].update(next_enabled=True),
            "bool page index": lambda v: v["pages"][0].update(page_index=False),
            "bool next": lambda v: v["pages"][0].update(next_enabled=0),
            "same-page duplicate": lambda v: v["pages"][0]["cards"].append(deepcopy(card(v))),
            "price null": lambda v: card(v).update(price_php=None),
            "price bool": lambda v: card(v).update(price_php=True),
            "price nan": lambda v: card(v).update(price_php=float("nan")),
            "price huge": lambda v: card(v).update(price_php=10**1000),
            "price negative": lambda v: card(v).update(price_php=-1),
            "title absent": lambda v: card(v).update(title=""),
            "wrong shop": lambda v: card(v).update(product_url="https://shopee.ph/product/1/101"),
            "wrong region": lambda v: card(v).update(product_url="https://shopee.sg/product/64922227/101"),
            "credentials": lambda v: card(v).update(product_url="https://user:secret@shopee.ph/product/64922227/101"),
            "http": lambda v: card(v).update(product_url="http://shopee.ph/product/64922227/101"),
            "empty models": lambda v: detail(v).update(skus=[]),
            "incomplete models": lambda v: detail(v).update(models_complete=False),
            "nonbool model flag": lambda v: detail(v).update(models_complete=1),
            "incomplete gallery": lambda v: detail(v).update(gallery_complete=False),
            "duplicate model": lambda v: detail(v)["skus"].append(deepcopy(detail(v)["skus"][0])),
            "numeric model": lambda v: detail(v)["skus"][0].update(model_id=1001),
            "unicode model": lambda v: detail(v)["skus"][0].update(model_id="１２３"),
            "empty props": lambda v: detail(v)["skus"][0].update(sku_props={}),
            "empty option": lambda v: detail(v)["skus"][0].update(sku_props={"颜色": ""}),
            "wrong image host": lambda v: detail(v).update(main_image_url="https://example.invalid/file/a"),
            "nonimage CDN path": lambda v: detail(v).update(main_image_url="https://down-ph.img.susercontent.com/review/a"),
            "relative image": lambda v: detail(v).update(main_image_url="/file/abc"),
            "thumbnail": lambda v: detail(v).update(main_image_url=image_url(1) + "_tn"),
            "duplicate gallery": lambda v: detail(v)["secondary_image_urls"].append(detail(v)["main_image_url"]),
            "bad sku image": lambda v: detail(v)["skus"][0].update(sku_image_url="javascript:alert(1)"),
        }
        for label, mutate in cases.items():
            with self.subTest(label=label):
                value = fixture()
                mutate(value)
                with self.assertRaises(ValueError):
                    publisher._convert(value)

    def test_incomplete_intermediate_page_and_repeated_page_rejected(self):
        value = fixture(31)
        value["pages"][0]["cards"].pop()
        with self.assertRaisesRegex(ValueError, "30"):
            publisher._convert(value)
        value = fixture(31)
        value["pages"][1]["cards"] = deepcopy(value["pages"][0]["cards"])
        with self.assertRaisesRegex(ValueError, "重复页"):
            publisher._convert(value)


class PublishTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="ugreen-publish-offline-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve() / "project"
        self.root.mkdir()
        self.daily = self.root / "worktrees/20000102_ugreen_topsales"
        self.root_spec = self.root / publisher.TEMPLATE_RELATIVE
        self.root_spec.parent.mkdir(parents=True)
        template(self.root_spec)
        (self.root / ".gitignore").write_text("/result/*\n/worktrees/\n__pycache__/\n", encoding="utf-8")
        source = self.root / publisher.SCRIPT_RELATIVE
        source.parent.mkdir(parents=True)
        shutil.copyfile(SCRIPTS / "publish_result.py", source)
        self.git(self.root, "init", "-b", "main")
        self.git(self.root, "add", ".")
        self.git(self.root, "-c", "user.name=Offline Fixture", "-c", "user.email=offline@example.invalid", "commit", "-m", "offline fixture")
        self.git(self.root, "worktree", "add", "-b", "daily/20000102-ugreen-topsales", str(self.daily), "HEAD")
        self.script_patch = patch.object(publisher, "__file__", str(self.daily / publisher.SCRIPT_RELATIVE))
        self.script_patch.start()
        self.addCleanup(self.script_patch.stop)
        self.output = self.daily / "result/ugreen_topsales.xlsx"
        self.output.parent.mkdir()
        self.output.write_bytes(b"previous offline workbook must survive failure")
        self.previous = self.output.read_bytes()
        self.template_hash = hashlib.sha256(self.root_spec.read_bytes()).hexdigest()

    @staticmethod
    def git(root, *args):
        return subprocess.run(["git", *args], cwd=root, check=True, text=True, capture_output=True).stdout

    def publish(self, value=None):
        return publisher.publish(fixture() if value is None else value, self.root, self.daily)

    def unchanged(self):
        self.assertEqual(self.output.read_bytes(), self.previous)
        self.assertEqual(sorted(path.name for path in self.output.parent.iterdir()), [self.output.name])
        self.assertEqual(hashlib.sha256(self.root_spec.read_bytes()).hexdigest(), self.template_hash)

    def test_full_roundtrip_cross_midnight_and_only_final_workbook_remains(self):
        receipt = self.publish()
        self.assertEqual(receipt["path"], str(self.output))
        self.assertEqual(receipt["sha256"], hashlib.sha256(self.output.read_bytes()).hexdigest())
        self.assertEqual(receipt["template_sha256"], self.template_hash)
        self.assertEqual(receipt["product_count"], 2)
        self.assertEqual(receipt["sku_count"], 2)
        self.assertEqual(receipt["main_image_count"], 2)
        self.assertEqual(receipt["secondary_image_count"], 2)
        self.assertEqual(receipt["audit_status"], "通过")
        self.assertEqual(excel.validate_excel(self.output).product_count, 2)
        wb = load_workbook(self.output)
        try:
            self.assertEqual(wb["商品汇总"]["B3"].value, "2000/01/03 00:01:00")
            self.assertEqual(wb["商品汇总"]["D7"].number_format, "@")
            self.assertEqual(wb["商品汇总"]["I7"].value, 1200)
            self.assertIsNone(wb["商品汇总"]["I8"].value)
            self.assertIsNone(wb["SKU明细"]["H7"].value)
        finally:
            wb.close()
        self.assertEqual(list(self.output.parent.iterdir()), [self.output])
        self.assertEqual(self.git(self.root, "status", "--porcelain"), "")
        self.assertEqual(self.git(self.daily, "status", "--porcelain"), "")
        self.assertEqual(hashlib.sha256(self.root_spec.read_bytes()).hexdigest(), self.template_hash)

    def test_multi_page_real_independent_validation(self):
        value = fixture(31)
        value["pages"][1]["cards"].insert(0, deepcopy(value["pages"][0]["cards"][0]))
        receipt = self.publish(value)
        self.assertEqual(receipt["list_page_count"], 2)
        self.assertEqual(receipt["product_count"], 31)

    def test_template_contract_reads_only_schema(self):
        receipt = publisher.validate_template_contract(self.root)
        self.assertEqual(receipt["sheets"], list(excel.SHEET_NAMES))
        self.assertEqual(receipt["header_row"], 6)
        self.assertEqual(receipt["sha256"], self.template_hash)
        self.unchanged()

    def test_bad_template_schema_rejected(self):
        workbook = load_workbook(self.root_spec)
        workbook["商品汇总"]["A6"] = "错误表头"
        workbook.save(self.root_spec)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "表头"):
            publisher.validate_template_contract(self.root)

    def test_wrong_date_is_rejected_and_no_today_dependency(self):
        value = fixture()
        value["started_at"] = "2000-01-01T10:00:00Z"
        with self.assertRaisesRegex(ValueError, "启动日"):
            self.publish(value)
        self.unchanged()

    def test_source_must_execute_inside_target_worktree(self):
        with patch.object(publisher, "__file__", str(self.root / publisher.SCRIPT_RELATIVE)):
            with self.assertRaisesRegex(ValueError, "代码必须"):
                self.publish()
        self.unchanged()

    def test_dirty_main_or_daily_rejected(self):
        for root in (self.root, self.daily):
            extra = root / "offline-untracked.txt"
            extra.write_text("fixture", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "未提交改动"):
                self.publish()
            extra.unlink()
            self.unchanged()

    def test_head_mismatch_rejected(self):
        self.git(self.daily, "-c", "user.name=Offline Fixture", "-c", "user.email=offline@example.invalid", "commit", "--allow-empty", "-m", "different offline HEAD")
        with self.assertRaisesRegex(ValueError, "同一已提交 HEAD"):
            self.publish()
        self.unchanged()

    def test_wrong_branch_rejected(self):
        self.git(self.daily, "branch", "-m", "daily/wrong-offline-branch")
        with self.assertRaisesRegex(ValueError, "分支"):
            self.publish()
        self.unchanged()

    def test_unknown_result_file_not_removed(self):
        unknown = self.output.parent / ".ugreen-publish-not-ours.xlsx"
        unknown.write_bytes(b"must not be ignored by a glob")
        with self.assertRaisesRegex(ValueError, "未知文件"):
            self.publish()
        self.assertEqual(unknown.read_bytes(), b"must not be ignored by a glob")
        self.assertEqual(self.output.read_bytes(), self.previous)

    def test_symlink_output_and_template_rejected(self):
        self.output.unlink()
        self.output.symlink_to(self.root_spec)
        with self.assertRaisesRegex(ValueError, "符号链接"):
            self.publish()
        self.assertEqual(hashlib.sha256(self.root_spec.read_bytes()).hexdigest(), self.template_hash)
        self.output.unlink()
        self.output.write_bytes(self.previous)
        daily_spec = self.daily / publisher.TEMPLATE_RELATIVE
        daily_spec.unlink()
        daily_spec.symlink_to(self.root_spec)
        with self.assertRaisesRegex(ValueError, "符号链接"):
            publisher.validate_template_contract(self.daily)
        self.unchanged()

    def test_parent_traversal_or_unregistered_directory_rejected(self):
        with self.assertRaisesRegex(ValueError, r"\.\."):
            publisher.publish(fixture(), self.root, self.daily / "../20000102_ugreen_topsales")
        other = self.root / "worktrees/20000103_ugreen_topsales"
        other.mkdir()
        with self.assertRaises(ValueError):
            publisher.publish(fixture(), self.root, other)
        self.unchanged()

    def test_repo_lock_prevents_concurrent_publication(self):
        lock = self.root / ".git/ugreen_daily_run.lock"
        with lock.open("a+") as handle:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            with self.assertRaisesRegex(ValueError, "另一个"):
                self.publish()
        self.unchanged()

    def test_independent_validation_failure_preserves_old_workbook(self):
        with patch.object(excel, "validate_excel", side_effect=ValueError("offline independent failure")):
            with self.assertRaisesRegex(ValueError, "independent failure"):
                self.publish()
        self.unchanged()

    def test_writer_failure_preserves_old_workbook(self):
        with patch.object(excel, "write_excel", side_effect=OSError("offline write failure")):
            with self.assertRaisesRegex(OSError, "write failure"):
                self.publish()
        self.unchanged()

    def test_mutation_before_final_replace_is_rechecked(self):
        real_validate = excel.validate_excel

        def validate_and_change(path):
            stats = real_validate(path)
            (self.root / "new-untracked.txt").write_text("changed during publication", encoding="utf-8")
            return stats

        with patch.object(excel, "validate_excel", side_effect=validate_and_change):
            with self.assertRaisesRegex(ValueError, "未提交改动"):
                self.publish()
        self.unchanged()

    def test_cli_prints_receipt_only_on_success(self):
        out, err = io.StringIO(), io.StringIO()
        with patch.object(sys, "stdin", io.StringIO(json.dumps(fixture()))), redirect_stdout(out), redirect_stderr(err):
            code = publisher.main(["--project-root", str(self.root), "--worktree", str(self.daily)])
        self.assertEqual(code, 0, err.getvalue())
        self.assertEqual(json.loads(out.getvalue())["audit_status"], "通过")
        self.assertEqual(err.getvalue(), "")

    def test_cli_rejects_duplicate_json_keys_without_writing(self):
        out, err = io.StringIO(), io.StringIO()
        with patch.object(sys, "stdin", io.StringIO('{"pages": [], "pages": []}')), redirect_stdout(out), redirect_stderr(err):
            code = publisher.main(["--project-root", str(self.root), "--worktree", str(self.daily)])
        self.assertEqual(code, 2)
        self.assertEqual(out.getvalue(), "")
        self.assertIn("重复键", err.getvalue())
        self.unchanged()


if __name__ == "__main__":
    unittest.main()
