"""Validated cold-start recovery of this run's cumulative workbook, never a crawl."""
from collections import defaultdict
import hashlib
from pathlib import Path
import re
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

import excel
from checkpoint_result import PROGRESS_NAME, HEADERS, build_snapshot, validate_snapshot
from checkpoint_state import read_embedded
from publish_result import _identity, _safe_path, _timestamp
from validate_result import _target_workbook

LEGACY_WARNINGS = [
    '旧版仅恢复已保存业务单元格；已完成商品的原列表标题未保存，内部沿用表中详情标题。',
    '旧版列表页 URL 按保存的店铺 URL 与页码规范化恢复，不声称原始 URL 拼写相同。',
    '旧版规格仅按展示分隔符恢复并核对往返一致；不保证还原未保存的原始字段边界。',
]


def _require(condition, message):
    if not condition:
        raise ValueError(message)


def pending_cards(payload):
    """Identity-based difference, not a count offset or exact URL comparison."""
    done = {_identity(detail['product_url']) for detail in payload['details']}
    seen, pending = set(), []
    for page in payload['pages']:
        for card in page['cards']:
            key = _identity(card['product_url'])
            if key not in seen and key not in done:
                pending.append(card)
            seen.add(key)
    return pending


def _legacy_payload(path):
    """Explicit migration only. The caller must bind the approved file hash."""
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        _require(tuple(wb.sheetnames) == excel.SHEET_NAMES, '进度表四表名称不符')
        tables = []
        for name, headers in zip(excel.SHEET_NAMES, HEADERS):
            sheet = wb[name]
            _require(sheet['A2'].value == f'UGREEN Top Sales {name}（采集进度）', '非本轮累计进度表')
            _require([c.value for c in sheet[6]] == headers, '进度表表头不符')
            tables.append([list(row) for row in sheet.iter_rows(min_row=7, values_only=True)
                           if any(value is not None for value in row)])
        summary, sku_rows, image_rows, audit_rows = tables
        audit = {row[0]: row[1] for row in audit_rows}
        _require(audit.get('跨页重复次数') == 0, '旧版表未保存跨页重复卡，不能无损恢复分页')
        _require(audit.get('商品卡出现次数') == audit.get('唯一商品数') == len(summary), '旧版列表计数不一致')
        match = re.fullmatch(r'已取得 (\d+)/(\d+) 页', str(audit_rows[2][2]))
        _require(match is not None, '已采集/总页数未明确保存')
        acquired, total = map(int, match.groups())
        _require(acquired == audit.get('列表总页数'), '已保存列表页数不一致')
        started = str(wb['商品汇总']['A4'].value).removeprefix('启动时间：')
        captured = str(wb['商品汇总']['A5'].value).removeprefix('保存时间：')
        for name in excel.SHEET_NAMES:
            _require(wb[name]['A4'].value == '启动时间：' + started, '各表启动时间不一致')
            _require(wb[name]['A5'].value == '保存时间：' + captured, '各表保存时间不一致')
        payload = dict(schema_version='1.0', store_url=audit['店铺 URL'],
                       started_at=started, captured_at=captured, pages=[], details=[])
        source_url = urlsplit(payload['store_url'])
        for index in range(acquired):
            query = dict(parse_qsl(source_url.query)); query['page'] = str(index)
            payload['pages'].append(dict(page_index=index, total_pages=total,
                url=urlunsplit(source_url._replace(query=urlencode(query))),
                next_enabled=index < total - 1, cards=[]))
        sku_groups, image_groups = defaultdict(list), defaultdict(list)
        for row in sku_rows:
            sku_groups[(row[2], row[3])].append(row)
        for row in image_rows:
            image_groups[(row[2], row[3])].append(row)
        for rank, row in enumerate(summary, 1):
            _require(row[0] == rank and type(row[1]) is int and 1 <= row[1] <= acquired, '旧版排名或页码不连续')
            page = payload['pages'][row[1] - 1]
            _require(row[2] == len(page['cards']) + 1, '旧版页内位置存在缺口')
            page['cards'].append(dict(product_url=row[13], title=row[5],
                                     price_php=row[6], monthly_sales_text=row[7]))
            key = (row[3], row[4])
            if row[14] != '已完成详情':
                _require(row[14] == '未采集详情（标题来自列表）', '未知详情采集状态')
                _require(not sku_groups[key] and not image_groups[key], '未完成商品含未知详情行')
                continue
            models = []
            for sku in sku_groups[key]:
                _require(isinstance(sku[6], str), '旧版规格不是文本')
                pairs = [part.split(': ', 1) for part in sku[6].split(', ')]
                _require(all(len(pair) == 2 and all(pair) for pair in pairs), '旧版规格不能安全拆解')
                props = dict(pairs)
                _require(len(props) == len(pairs), '旧版规格键重复，不能安全还原')
                _require(', '.join(f'{k}: {v}' for k, v in props.items()) == sku[6], '旧版规格往返不一致')
                models.append(dict(model_id=sku[5], sku_props=props, sku_image_url=sku[7]))
            images = image_groups[key]
            _require(images and images[0][5:8] == ['主图', 1, row[11]], '旧版主图不一致')
            _require([r[5:7] for r in images[1:]] == [['副图', n] for n in range(1, len(images))], '副图顺序不连续')
            _require(len(models) == row[10] and len(images) - 1 == row[12], 'SKU/图库计数不一致')
            payload['details'].append(dict(product_url=row[13], title=row[5], skus=models,
                main_image_url=row[11], secondary_image_urls=[r[7] for r in images[1:]],
                models_complete=True, gallery_complete=True))
        snapshot = build_snapshot(payload)
        _require(snapshot['sheets'] == tables, '恢复后与旧表业务单元格不一致')
        return payload
    finally:
        wb.close()


def read_checkpoint_file(path, *, legacy_sha256=None):
    """Read-only roundtrip verification; corruption never falls back to legacy."""
    path = _safe_path(Path(path))
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if legacy_sha256 is not None:
        _require(re.fullmatch(r'[0-9a-f]{64}', legacy_sha256) is not None and legacy_sha256 == digest,
                 '指定的旧版进度表 SHA-256 不匹配')
    state = read_embedded(path)
    if state is None:
        _require(legacy_sha256 is not None, '旧版进度表无恢复状态；需明确审核并绑定 SHA-256 后迁移')
        state = dict(payload=_legacy_payload(path), origin=dict(kind='legacy_visible_rows',
                     source_sha256=digest, warnings=LEGACY_WARNINGS))
    payload = state['payload']
    snapshot = build_snapshot(payload)
    validate_snapshot(path, snapshot)
    _require(hashlib.sha256(path.read_bytes()).hexdigest() == digest, '读取期间进度表发生变化')
    pending = pending_cards(payload)
    _require(len(pending) == snapshot['counts']['pending_detail_count'], '恢复待办数量不一致')
    done = snapshot['counts']['detail_count']
    return dict(path=str(path), sha256=digest, payload=payload, origin=state['origin'],
                counts=snapshot['counts'], pending_cards=pending, saved_count=done,
                next_checkpoint_at=(done // 50 + 1) * 50)


def load_checkpoint(project_root, worktree, *, legacy_sha256=None):
    root, daily = _safe_path(Path(project_root)), _safe_path(Path(worktree))
    target = _target_workbook(root, daily).with_name(PROGRESS_NAME)
    recovered = read_checkpoint_file(target, legacy_sha256=legacy_sha256)
    started = _timestamp(recovered['payload']['started_at'], 'started_at')
    _require(daily.name == started.astimezone(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d_ugreen_topsales'),
             '恢复状态启动日不属于目标 worktree')
    return recovered
