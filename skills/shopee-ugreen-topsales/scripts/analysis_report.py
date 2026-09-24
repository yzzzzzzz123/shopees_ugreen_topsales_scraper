"""由已验收的每日完整 Excel 生成可复算的 PDF；不访问浏览器。"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
import hashlib
import html
import os
from pathlib import Path
import re
import tempfile
from urllib.parse import quote

from openpyxl import load_workbook
from excel import validate_excel
from validate_result import _target_workbook

REPORT_NAME = "ugreen_topsales_analysis.pdf"
REPORT_MARKER = "<!-- ugreen-daily-analysis:v1 -->"
PRICE_BANDS = ((0, 100), (100, 300), (300, 500), (500, 1000),
               (1000, 2000), (2000, 5000), (5000, None))
MISSING = "缺数据"


@dataclass(frozen=True)
class Product:
    shop_id: str
    item_id: str
    rank: int
    title: str
    url: str
    price: Decimal
    monthly: int | None
    monthly_text: str | None

    @property
    def key(self) -> tuple[str, str]:
        return self.shop_id, self.item_id


@dataclass
class Snapshot:
    day: date
    path: Path
    sha256: str
    captured_at: str
    products: dict[tuple[str, str], Product]
    revision: str | None = None


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_snapshot(root: Path, daily: Path) -> Snapshot:
    path = _target_workbook(root, daily)
    return _load_snapshot_file(path, daily)


def _load_snapshot_file(path: Path, daily: Path) -> Snapshot:
    before = digest(path)
    stats = validate_excel(path)
    if stats.collection_mode != "full_topsales" or stats.audit_status != "通过":
        raise ValueError("分析只接受验收通过的完整 Top Sales 正式表")
    day = datetime.strptime(daily.name[:8], "%Y%m%d").date()
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["商品汇总"]
        captured = sheet["B3"].value
        captured_day = datetime.strptime(captured, "%Y/%m/%d %H:%M:%S").date()
        if captured_day < day:
            raise ValueError("工作簿采集时间早于日期工作区业务日")
        products = {}
        for row in sheet.iter_rows(min_row=7, max_col=14, values_only=True):
            p = Product(str(int(row[3])), str(int(row[4])), row[0], row[5], row[13],
                        Decimal(str(row[6])), row[8], row[7])
            if p.key in products:
                raise ValueError("分析输入出现重复商品身份")
            products[p.key] = p
    finally:
        workbook.close()
    if before != digest(path):
        raise ValueError("分析读取期间源工作簿发生变化")
    return Snapshot(day, path, before, captured, products)


def load_history(root: Path, daily: Path) -> tuple[Snapshot, dict[date, Snapshot], list[str]]:
    current = load_snapshot(root, daily)
    history = {current.day: current}
    excluded = []
    for folder in sorted((root / "worktrees").iterdir()):
        if not re.fullmatch(r"[0-9]{8}_ugreen_topsales", folder.name) or folder == daily:
            continue
        try:
            day = datetime.strptime(folder.name[:8], "%Y%m%d").date()
        except ValueError:
            continue
        if day >= current.day:
            continue
        try:
            from snapshot_history import latest_snapshot
            issues = []
            snapshot = latest_snapshot(root, folder, current.captured_at, issues)
            excluded.extend(f"{day.isoformat()}：{issue}" for issue in issues)
            if snapshot is not None:
                history[day] = snapshot
        except Exception as exc:
            excluded.append(f"{day.isoformat()}：{exc}")
    return current, history, excluded


def totals(snapshot: Snapshot) -> dict:
    products = list(snapshot.products.values())
    known = [p for p in products if p.monthly is not None]
    volume = sum(p.monthly for p in known) if known else None
    amount = sum((p.price * p.monthly for p in known), Decimal(0)) if known else None
    active = sum(p.monthly > 0 for p in known)
    return dict(links=len(products), known=len(known), missing=len(products)-len(known),
                active=active, zero=sum(p.monthly == 0 for p in known), volume=volume,
                amount=amount, weighted_price=amount / volume if volume else None,
                rounded=sum(bool(re.search(r"[KkMm+]", re.sub(r"\s*Sold/Month\s*$", "", p.monthly_text or "", flags=re.I))) for p in known))


def delta(now, before):
    if now is None or before is None:
        return None, None
    difference = Decimal(str(now)) - Decimal(str(before))
    percent = difference / Decimal(str(before)) * 100 if before != 0 else None
    return difference, percent


def movement(current: Snapshot, previous: Snapshot | None, field: str) -> dict:
    if previous is None:
        return dict(up=None, down=None, flat=None, unknown=None, common=None)
    result = dict(up=0, down=0, flat=0, unknown=0, common=0)
    for key in current.products.keys() & previous.products.keys():
        result["common"] += 1
        now, before = getattr(current.products[key], field), getattr(previous.products[key], field)
        if now is None or before is None:
            result["unknown"] += 1
        else:
            result["up" if now > before else "down" if now < before else "flat"] += 1
    return result


def listing_changes(day: date, history: dict[date, Snapshot]) -> dict:
    now, previous = history.get(day), history.get(day - timedelta(days=1))
    if now is None or previous is None:
        return dict(added=None, first_seen=None, returned=None, missing=None, added_keys=set())
    added = now.products.keys() - previous.products.keys()
    seen = set().union(*(set(s.products) for d, s in history.items() if d < day))
    return dict(added=len(added), first_seen=len(added-seen), returned=len(added & seen),
                missing=len(previous.products.keys()-now.products.keys()), added_keys=added)


def price_bands(snapshot: Snapshot) -> list[dict]:
    total = totals(snapshot)
    result = []
    for low, high in PRICE_BANDS:
        selected = [p for p in snapshot.products.values() if p.price >= low and (high is None or p.price < high)]
        known = [p for p in selected if p.monthly is not None]
        volume = sum(p.monthly for p in known) if known else (0 if not selected else None)
        result.append(dict(label=f"[{low:,}, {high:,})" if high is not None else f"≥{low:,}",
                           links=len(selected), active=sum(p.monthly > 0 for p in known),
                           unknown=len(selected)-len(known), volume=volume,
                           share=Decimal(volume)/total["volume"]*100 if total["volume"] and volume is not None else None,
                           amount=sum((p.price*p.monthly for p in known), Decimal(0)) if known or not selected else None))
    return result


def number(value, decimals=0, signed=False) -> str:
    if value is None:
        return MISSING
    prefix = "+" if signed and value > 0 else ""
    return prefix + f"{value:,.{decimals}f}"


def percentage(numerator, denominator) -> str:
    return "不可计算" if not denominator else f"{Decimal(numerator)/Decimal(denominator)*100:.2f}%"


def percent_change(now, before) -> str:
    change, percent = delta(now, before)
    if change is None:
        return MISSING
    if percent is None:
        return "基期为0，比例不适用"
    return number(percent, 2, signed=True) + "%"


def escape(value) -> str:
    return html.escape(str(value), quote=False).replace("|", "&#124;").replace("`", "&#96;").replace("\n", " ").replace("\r", " ")


def link(product: Product, short=False) -> str:
    title = product.title[:60] + ("…" if len(product.title) > 60 else "") if short else product.title
    safe_url = quote(product.url, safe=":/?=&%._~!$'()*+,-;@")
    # HTML links avoid Markdown bracket injection from product titles.
    return f'<a href="{html.escape(safe_url, quote=True)}">{product.item_id} · {escape(title)}</a>'


def table(headers, rows) -> str:
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"] +
                     ["| " + " | ".join(map(str, row)) + " |" for row in rows])


def days_ending(day: date, count: int) -> list[date]:
    return [day - timedelta(days=n) for n in range(count-1, -1, -1)]


def sparkline(values) -> str:
    known = [v for v in values if v is not None]
    maximum = max(known, default=0)
    bars = "▁▂▃▄▅▆▇█"
    return "".join("·" if v is None else bars[round(v/maximum*7)] if maximum else bars[0] for v in values)


def trend_chart(days: list[date], values, label: str) -> str:
    """Markdown 内的等间隔点趋势图；缺测不补零、不跨缺口连线。"""
    known = [v for v in values if v is not None]
    if not known:
        return (f"{label}：**没有可计算的日观测值，暂无法绘制数值趋势。**\n\n"
                f"窗口 {days[0]} 至 {days[-1]}，每天一格：`{sparkline(values)}`。"
                "`·` 表示缺数据，不表示 0。")
    maximum = max(known, default=0)
    height = 6
    points = [None if v is None else round(v/maximum*height) if maximum else 0 for v in values]
    lines = ["```text", label]
    for level in range(height, -1, -1):
        value = maximum*level/height if maximum else 0
        cells = " ".join("●" if p == level else " " for p in points)
        lines.append(f"{value:>10,.0f} ┤ {cells}")
    lines += ["           └" + "──"*len(days),
              "             " + " ".join(str((i+1)//10 or " ") for i in range(len(days))),
              "             " + " ".join(str((i+1)%10) for i in range(len(days))),
              "数据状态：" + " ".join("·" if v is None else "有" for v in values), "```",
              f"横轴 01={days[0]}，{len(days):02d}={days[-1]}，每天一格；● 为观测点，· 为缺数据。"]
    if len(known) < 2:
        lines.append("**有效观测不足 2 天，尚不能判断趋势；未用零值或插值补齐。**")
    return "\n".join(lines)


def render_report(current: Snapshot, history: dict[date, Snapshot], excluded: list[str]) -> str:
    day = current.day
    previous = history.get(day-timedelta(days=1))
    current_total = totals(current)
    previous_total = totals(previous) if previous else None
    monthly = movement(current, previous, "monthly")
    price = movement(current, previous, "price")
    listings = listing_changes(day, history)
    common = current.products.keys() & previous.products.keys() if previous else set()
    comparable = [key for key in common if current.products[key].monthly is not None and previous.products[key].monthly is not None]
    now_common = sum(current.products[k].monthly for k in comparable) if comparable else None
    before_common = sum(previous.products[k].monthly for k in comparable) if comparable else None
    sources = sorted(history.items())
    lines = [REPORT_MARKER, f"# UGREEN 每日经营分析（{day}）", "",
             f"业务日：**{day}**（上海启动日）。本轮采集完成时间：**{current.captured_at} Asia/Shanghai**。",
             "国家 / 币种：PH / PHP。店铺：UGREEN（64922227）。本报告只读正式 Excel，未重新访问网页。", "",
             "## 数据覆盖与阅读口径", "",
             f"- 已验收完整快照：**{len(history)} 天**；最近 7 天有 **{sum(d in history for d in days_ending(day,7))}/7** 天，最近 30 天有 **{sum(d in history for d in days_ending(day,30))}/30** 天。",
             f"- 昨日（{day-timedelta(days=1)}）快照：**{'可用' if previous else '缺失，所有日环比及新增链接数量暂不可计算'}**。",
             f"- 本日带 K/M/+ 展示精度限制的月销链接：{current_total['rounded']} 件；未展示月销：{current_total['missing']} 件。",
             "- “月销”是每日观察到的滚动月销展示值/下界，不是当日订单；两日差值不是新增订单或净销量。K/M/+ 的展示精度不能恢复成精确销量。",
             "- 估算月销额 = Σ（当前列表价 × 月销展示下界），仅覆盖有月销的商品；受价格变动、SKU 价差和促销影响，既不是真实 GMV，也不保证是真实销售额下界。",
             "- 每个产品的“客单价”只以当前列表价作单件价格代理。真实订单客单价需要订单金额和订单数，现有快照无法推得。店铺加权单件价格 = 估算月销额 / 可观测月销合计。",
             "- 新出现链接 = 今日身份集合减昨日集合，是上新的观察代理；首次观察、重新出现分别列出，不直接断言真实上架或下架时间。首份基线不把全店链接计为新增。",
             "- 缺采、未展示月销与真实 0 分开；缺昨天不拿更早一天冒充昨天。7/30 日均按截至业务日的自然日窗口，缺口不补零、不插值。", "",
             "## 店铺汇总", ""]
    volume_change = delta(current_total["volume"], previous_total["volume"] if previous_total else None)[0]
    lines.append(table(["指标", "本日结果", "口径 / 对比"], [
        ["商品链接数", number(current_total["links"]), "全部唯一 shop_id + item_id"],
        ["可观测月销合计", number(current_total["volume"]), f"覆盖 {current_total['known']} 件；{current_total['missing']} 件未展示"],
        ["店铺可观测月销合计日变化", number(volume_change, signed=True), "商品集合与展示覆盖可能变化，不等于同款增长"],
        ["同款可比月销变化", number(delta(now_common,before_common)[0], signed=True), f"可比链接数 {len(comparable) if previous else MISSING}；增幅 {percent_change(now_common,before_common)}"],
        ["估算月销额（PHP）", number(current_total["amount"],2), "当前价 × 月销下界，仅已展示样本"],
        ["月销加权单件价格（PHP）", number(current_total["weighted_price"],2), "金额代理 / 月销合计；非订单客单价"],
        ["月销上涨 / 下滑 / 持平链接", f"{number(monthly['up'])} / {number(monthly['down'])} / {number(monthly['flat'])}", f"共同链接 {number(monthly['common'])}；月销不可比 {number(monthly['unknown'])}"],
        ["价格上涨 / 下调 / 持平链接", f"{number(price['up'])} / {number(price['down'])} / {number(price['flat'])}", f"共同链接 {number(price['common'])}；价格按 PHP 比较"],
        ["新出现 / 首次观察 / 重新出现", f"{number(listings['added'])} / {number(listings['first_seen'])} / {number(listings['returned'])}", "均以昨日完整集合为前提，首次观察只对已保存历史负责"],
        ["昨日有、今日未观察到", number(listings["missing"]), "不能仅凭列表消失断言下架"],
        ["有月销 / 零月销 / 未知链接", f"{current_total['active']} / {current_total['zero']} / {current_total['missing']}", "滚动月销 >0 / =0 / 未展示"],
        ["可确认动销链接占全店", percentage(current_total["active"],current_total["links"]), "未知保留在全店分母中，保守观察比例"],
        ["已知月销样本动销率", percentage(current_total["active"],current_total["known"]), "仅已展示月销链接为分母"],
    ]))
    lines += ["", "## 固定价格段的出单占比与动销", "",
              "价格区间左闭右开（PHP），按当日列表价归档。“出单占比”使用该段可观测月销 / 全店可观测月销作代理，不代表当天订单占比；未知月销不进入销量分母。", "",
              table(["价格段（PHP）", "链接数", "有月销链接", "月销未知", "月销合计", "销量占比", "估算月销额（PHP）"],
                    [[b["label"],b["links"],b["active"],b["unknown"],number(b["volume"]),
                      number(b["share"],2)+"%" if b["share"] is not None else "不可计算",number(b["amount"],2)] for b in price_bands(current)]), ""]
    for count in (7,30):
        dates = days_ending(day,count)
        values = [totals(history[d])["volume"] if d in history else None for d in dates]
        lines += [f"## {count} 日店铺月销趋势图", "", trend_chart(dates, values, "店铺可观测月销合计（展示下界，件）"), ""]
    dates = days_ending(day,30)
    additions = [listing_changes(d,history)["added"] for d in dates]
    lines += ["## 30 日新出现链接趋势图（上新代理）", "", trend_chart(dates, additions, "与前一自然日相比的新出现链接数"), "",
              "### 逐日数据与覆盖", "",
              table(["业务日", "链接数", "已知月销链接", "月销合计", "估算月销额（PHP）", "新出现", "首次观察", "重新出现"],
                    [[d.isoformat(), number(totals(history[d])["links"]) if d in history else MISSING,
                      number(totals(history[d])["known"]) if d in history else MISSING,
                      number(totals(history[d])["volume"]) if d in history else MISSING,
                      number(totals(history[d])["amount"],2) if d in history else MISSING,
                      number(listing_changes(d,history)["added"]),number(listing_changes(d,history)["first_seen"]),
                      number(listing_changes(d,history)["returned"])] for d in dates]), "",
              "## 全部产品链接的月销、价格变化与金额估算", "",
              "按本日 Top Sales 顺序。差值为本日减昨日；基期为 0 时报告绝对差值，百分比不计算。未展示与昨日缺失均不补零。", ""]
    product_rows = []
    trend_rows = []
    for product in sorted(current.products.values(),key=lambda p:p.rank):
        before = previous.products.get(product.key) if previous else None
        monthly_before = before.monthly if before else None
        price_before = before.price if before else None
        product_rows.append([product.rank,link(product,True),escape(product.monthly_text or "未展示"),number(monthly_before),
                             number(delta(product.monthly,monthly_before)[0],signed=True),percent_change(product.monthly,monthly_before),
                             number(product.price,2),number(price_before,2),number(delta(product.price,price_before)[0],2,True),
                             percent_change(product.price,price_before),number(product.price*product.monthly,2) if product.monthly is not None else MISSING])
        values = [history[d].products[product.key].monthly if d in history and product.key in history[d].products else None for d in dates]
        known = [v for v in values if v is not None]
        trend_rows.append([product.rank,link(product,True),f"`{sparkline(values[-7:])}`",f"`{sparkline(values)}`",
                           f"{sum(v is not None for v in values[-7:])}/7；{len(known)}/30",
                           f"{min(known):,}～{max(known):,}" if known else MISSING])
    lines += [table(["排名", "产品链接", "本日月销原文", "昨日月销", "月销差值", "月销变化率", "当前价/客单代理 PHP", "昨日价 PHP", "价格差值 PHP", "价格变化率", "估算月销额 PHP"],product_rows), "",
              "## 全部产品的 7 日 / 30 日月销迷你趋势图", "",
              f"7 日：{dates[-7]} 至 {day}；30 日：{dates[0]} 至 {day}。每个字符代表一天，从左到右按时间排列。",
              "`·` 是缺采、当日未观察到该商品或未展示月销，`▁` 至 `█` 是月销柱高；真实零为 `▁`。每个商品独立采用 0～其窗口最大值的纵轴，配合数值范围阅读，不跨商品比较柱高。单个观测点不构成趋势。", "",
              table(["排名", "产品链接", "7 日趋势", "30 日趋势", "有效天数", "30 日观测最小～最大"],trend_rows), "",
              "## 今日新出现链接明细", ""]
    if previous is None:
        lines.append("缺少昨日完整快照，本日仅建立观察基线；新增、首次观察和重新出现数量均不可判定。")
    elif not listings["added_keys"]:
        lines.append("与昨日完整快照相比，新出现链接为 0。")
    else:
        seen = set().union(*(set(s.products) for d,s in history.items() if d < day))
        lines.append(table(["产品链接", "观察分类", "当前价（PHP）", "月销原文"],
                           [[link(current.products[k]),"重新出现" if k in seen else "首次观察",number(current.products[k].price,2),
                             escape(current.products[k].monthly_text or "未展示")] for k in sorted(listings["added_keys"])]))
    lines += ["", "## 来源与核对", "",
              "仅使用当前店铺通过完整验收的正式快照；累计进度表、规范模板及不完整结果不进入日环比。没有其他竞品店铺的完整快照时，不生成跨店排名。", "",
              table(["业务日", "完成时间（上海）", "正式快照", "SHA-256"],
                    [[d.isoformat(),s.captured_at,f"`{s.path.parent.parent.name}/result/{s.path.name}"+(f" @ {s.revision[:12]}" if s.revision else "")+"`",f"`{s.sha256}`"] for d,s in sources]), "",
              f"核对：本日商品 {current_total['links']} 件；有月销 {current_total['known']} + 未展示 {current_total['missing']} = {current_total['links']}；",
              f"有月销 {current_total['active']} + 零月销 {current_total['zero']} = 已知样本 {current_total['known']}；所有价格段链接数、销量和金额分别回加到店铺合计。",
              "报告可由同一组源 Excel 确定性重算；修改源表、历史覆盖或报告正文后，独立报告验收会失败，须重新生成。"]
    if excluded:
        lines += ["", "排除的历史快照（未参与计算）："] + ["- "+escape(item) for item in excluded]
    return "\n".join(lines).rstrip()+"\n"


def _report_path(target: Path, legacy_sha256: str | None = None) -> Path:
    from analysis_pdf import check_controlled_pdf
    report = target.with_name(REPORT_NAME)
    check_controlled_pdf(report, legacy_sha256)
    return report


def _check_sources(history):
    from snapshot_history import read_blob
    for snapshot in history.values():
        actual=(hashlib.sha256(read_blob(snapshot.path.parent.parent,snapshot.revision)).hexdigest()
                if snapshot.revision else digest(snapshot.path))
        if actual != snapshot.sha256:
            raise ValueError("分析期间来源快照发生变化")


def write_report(root: Path, daily: Path, legacy_pdf_sha256: str | None = None) -> dict:
    from analysis_pdf import render_pdf
    from publish_result import _check_destination, _exclusive_publish
    target = _target_workbook(root,daily)
    root, daily = root.resolve(), daily.resolve()
    report = _report_path(target, legacy_pdf_sha256)
    with _exclusive_publish(root):
        _check_destination(target, legacy_pdf_sha256=legacy_pdf_sha256)
        current, history, excluded = load_history(root,daily)
        content = render_pdf(render_report(current,history,excluded))
        fd, name = tempfile.mkstemp(prefix=".ugreen-analysis-",suffix=".pdf",dir=target.parent)
        candidate = Path(name)
        try:
            with os.fdopen(fd,"wb") as stream:
                stream.write(content); stream.flush(); os.fsync(stream.fileno())
            if candidate.read_bytes() != content:
                raise ValueError("PDF 候选回读不一致")
            _check_sources(history)
            _report_path(target, legacy_pdf_sha256)
            os.replace(candidate,report)
        finally:
            candidate.unlink(missing_ok=True)
    return dict(path=str(report),sha256=digest(report),report_date=current.day.isoformat(),
                history_days=len(history),previous_day_available=current.day-timedelta(days=1) in history,
                product_count=len(current.products),analysis_status="通过",
                monthly_sales_lower_bound=totals(current)["volume"],
                estimated_monthly_amount_php=str(totals(current)["amount"]) if totals(current)["amount"] is not None else None)


def validate_report(root: Path, daily: Path) -> dict:
    from analysis_pdf import render_pdf
    current, history, excluded = load_history(root,daily)
    report = _report_path(current.path)
    if not report.is_file():
        raise ValueError("缺少每日 PDF 分析报告，请先执行 --analyze")
    if report.read_bytes() != render_pdf(render_report(current,history,excluded)):
        raise ValueError("每日分析报告与当前已验收快照不一致，请重新执行 --analyze")
    _check_sources(history)
    return dict(path=str(report),sha256=digest(report),analysis_status="通过",history_days=len(history))
