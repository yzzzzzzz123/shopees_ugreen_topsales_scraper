"""Cumulative partial Excel snapshots. Never connects to a browser."""
from __future__ import annotations
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from io import BytesIO
import zipfile
from xml.etree import ElementTree as ET
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
import excel
from publish_result import _check_destination, _identity, _image_url, _list_url, _monthly, _object, _text, _timestamp, _safe_path, _exclusive_publish
from validate_result import _target_workbook

PROGRESS_NAME = "ugreen_topsales_progress.xlsx"
MODE = "partial_topsales"
HEADERS = [list(excel.SUMMARY_HEADERS) + ["详情采集状态"], list(excel.SKU_HEADERS), list(excel.IMAGE_HEADERS), list(excel.AUDIT_HEADERS)]


def add_hyperlinks(path, snapshot):
    """Add standard OOXML links; the authoring API lacks a documented link setter."""
    sheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    doc_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    changed = {}
    with zipfile.ZipFile(path) as source:
        for index, rows in enumerate(snapshot["sheets"], 1):
            name = f"xl/worksheets/sheet{index}.xml"
            rel_name = f"xl/worksheets/_rels/sheet{index}.xml.rels"
            sheet = ET.fromstring(source.read(name))
            rels = ET.fromstring(source.read(rel_name)) if rel_name in source.namelist() else ET.Element(f"{{{rel_ns}}}Relationships")
            links = ET.Element(f"{{{sheet_ns}}}hyperlinks")
            for r, row in enumerate(rows, 7):
                for c, value in enumerate(row, 1):
                    if isinstance(value, str) and value.startswith("https://"):
                        rid = f"ugreenLink{r}_{c}"
                        ET.SubElement(links, f"{{{sheet_ns}}}hyperlink", {"ref": f"{chr(64+c)}{r}", f"{{{doc_rel}}}id": rid})
                        ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {"Id": rid, "Type": doc_rel + "/hyperlink", "Target": value, "TargetMode": "External"})
            if len(links):
                before = {"printOptions", "pageMargins", "pageSetup", "headerFooter", "drawing", "tableParts", "extLst"}
                pos = next((i for i, node in enumerate(sheet) if node.tag.split("}")[-1] in before), len(sheet))
                sheet.insert(pos, links)
                changed[name] = ET.tostring(sheet, encoding="utf-8", xml_declaration=True)
                changed[rel_name] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as output:
            for entry in source.infolist():
                output.writestr(entry, changed.pop(entry.filename, source.read(entry.filename)))
            for name, data in changed.items():
                output.writestr(name, data)
    path.write_bytes(buffer.getvalue())

def build_snapshot(payload):
    _object(payload, {"schema_version", "store_url", "started_at", "captured_at", "pages", "details"}, "进度载荷")
    if payload["schema_version"] != "1.0":
        raise ValueError("只支持 schema_version=1.0")
    _list_url(payload["store_url"], 0)
    started = _timestamp(payload["started_at"], "started_at")
    saved = _timestamp(payload["captured_at"], "captured_at")
    if saved < started:
        raise ValueError("保存时间不能早于启动时间")
    pages, details = payload["pages"], payload["details"]
    if type(pages) is not list or not pages or type(details) is not list:
        raise ValueError("需至少一页已验收列表，details 必须是数组")
    listed, occurrences = {}, 0
    total_pages = pages[0]["total_pages"]
    if type(total_pages) is not int or total_pages < len(pages):
        raise ValueError("列表总页数无效")
    for index, page in enumerate(pages):
        _object(page, {"page_index", "total_pages", "url", "next_enabled", "cards"}, "列表页")
        if type(page["page_index"]) is not int or page["page_index"] != index or type(page["total_pages"]) is not int or page["total_pages"] != total_pages:
            raise ValueError("已取得列表必须是连续前缀且总页数一致")
        _list_url(page["url"], index)
        if type(page["next_enabled"]) is not bool or page["next_enabled"] != (index < total_pages - 1):
            raise ValueError("分页按钮状态不符")
        cards = page["cards"]
        if type(cards) is not list or not 1 <= len(cards) <= 30 or (index < total_pages - 1 and len(cards) != 30):
            raise ValueError("列表商品卡数量不符")
        seen, new = set(), 0
        for position, card in enumerate(cards, 1):
            _object(card, {"product_url", "title", "price_php", "monthly_sales_text"}, "商品卡")
            key = _identity(card["product_url"])
            if key in seen:
                raise ValueError("同页商品身份重复")
            seen.add(key)
            _text(card["title"], "列表标题")
            price = card["price_php"]
            if type(price) not in (int, float) or not 0 <= price <= 2**53 - 1:
                raise ValueError("列表价格不是有效非负数值")
            monthly, lower, status = _monthly(card["monthly_sales_text"])
            occurrences += 1
            if key not in listed:
                new += 1
                listed[key] = [len(listed) + 1, index + 1, position, *key, card["title"], price,
                               monthly, lower, "已展示" if status == "displayed" else "未展示",
                               None, None, None, card["product_url"], "未采集详情（标题来自列表）"]
        if not new:
            raise ValueError("后续列表无新身份")
    verified = {}
    sku_rows, image_rows = [], []
    for detail in details:
        _object(detail, {"product_url", "title", "skus", "main_image_url", "secondary_image_urls", "models_complete", "gallery_complete"}, "详情")
        key = _identity(detail["product_url"])
        if key not in listed or key in verified:
            raise ValueError("详情重复或超出本轮列表")
        if detail["models_complete"] is not True or detail["gallery_complete"] is not True:
            raise ValueError("详情的 SKU/图库完整证据未确认")
        title = _text(detail["title"], "详情标题")
        skus = detail["skus"]
        if type(skus) is not list or not skus:
            raise ValueError("已完成详情须有真实 SKU")
        main = _image_url(detail["main_image_url"])
        if type(detail["secondary_image_urls"]) is not list:
            raise ValueError("副图必须为数组")
        secondary = [_image_url(url) for url in detail["secondary_image_urls"]]
        if len(set([main, *secondary])) != len(secondary) + 1:
            raise ValueError("主副图重复")
        row = listed[key]
        row[5], row[10], row[11], row[12], row[13], row[14] = title, len(skus), main, len(secondary), detail["product_url"], "已完成详情"
        common = [row[0], row[1], *key, title]
        models = set()
        for sku in skus:
            _object(sku, {"model_id", "sku_props", "sku_image_url"}, "SKU")
            model = sku["model_id"]
            if not isinstance(model, str) or not re.fullmatch("[0-9]+", model) or int(model) in models:
                raise ValueError("model_id 无效或重复")
            models.add(int(model))
            props = sku["sku_props"]
            if type(props) is not dict or not props:
                raise ValueError("SKU 规格不能为空")
            if len(", ".join(f"{k}: {v}" for k, v in props.items())) > 32767:
                raise ValueError("合并后的 SKU 规格超过 Excel 文本长度")
            for name, value in props.items():
                _text(name, "规格名")
                _text(value, "规格值")
            link = sku["sku_image_url"]
            if link is not None:
                _image_url(link)
            sku_rows.append(common + [model, ", ".join(f"{k}: {v}" for k, v in props.items()), link, detail["product_url"]])
        image_rows.append(common + ["主图", 1, main, detail["product_url"]])
        image_rows.extend(common + ["副图", n, url, detail["product_url"]] for n, url in enumerate(secondary, 1))
        verified[key] = detail
    sku_rows.sort(key=lambda row: row[0])
    image_rows.sort(key=lambda row: row[0])
    count, done = len(listed), len(verified)
    secondary_count = len(image_rows) - done
    status = f"部分采集：详情 {done}/{count}，待采集 {count-done}"
    audit = [
        ["店铺 URL", payload["store_url"], "本轮 PH UGREEN"],
        ["排序方式", "Top Sales", "按本轮列表首次出现顺序"],
        ["列表总页数", len(pages), f"已取得 {len(pages)}/{total_pages} 页"],
        ["商品卡出现次数", occurrences, "已取得分页卡片数"],
        ["跨页重复次数", occurrences-count, "按身份去重"],
        ["唯一商品数", count, "商品汇总含待采集详情的列表记录"],
        ["成功进入详情页", done, "仅计已完成并验收详情"],
        ["详情页失败数", None, "未统计；未采集不等于失败"],
        ["详情解析成功数", done, "SKU 与图库完整"],
        ["详情解析错误数", None, "未统计"],
        ["未展示月销商品数", sum(r[9] == "未展示" for r in listed.values()), "来自列表"],
        ["SKU 总数", len(sku_rows), "仅已完成详情"],
        ["有 SKU 的商品数", done, "未采集详情不补占位 SKU"],
        ["主图数", done, "仅已完成详情"],
        ["副图数", secondary_count, "仅已完成详情"],
        ["无效图片 URL 数", 0, "本次已保存图片 URL 均通过校验"],
        ["完整性结论", "未完成", status],
    ]
    counts = dict(product_count=count, detail_count=done, pending_detail_count=count-done,
                  sku_count=len(sku_rows), main_image_count=done, secondary_image_count=secondary_count,
                  list_page_count=len(pages), total_pages=total_pages)
    return dict(mode=MODE, started_at=payload["started_at"], saved_at=payload["captured_at"],
                status=status, counts=counts, headers=HEADERS,
                sheets=[list(listed.values()), sku_rows, image_rows, audit])

def validate_snapshot(path, snapshot):
    wb = load_workbook(path, read_only=False, data_only=False)
    try:
        if tuple(wb.sheetnames) != excel.SHEET_NAMES:
            raise ValueError("进度表四表名称不符")
        for name, headers, expected in zip(excel.SHEET_NAMES, HEADERS, snapshot["sheets"]):
            ws = wb[name]
            if ws.max_row != max(7, 6 + len(expected)):
                raise ValueError("进度表数据行数不符：" + name)
            if [c.value for c in ws[6]] != headers:
                raise ValueError("进度表表头不符：" + name)
            actual = [[c.value for c in row] for row in ws.iter_rows(min_row=7, max_row=6+len(expected))]
            if expected and actual != expected:
                mismatch = next(((r+7, c+1, repr(a)[:80], repr(b)[:80]) for r, (ra, rb) in enumerate(zip(actual, expected)) for c, (a, b) in enumerate(zip(ra, rb)) if a != b), (len(actual), len(expected)))
                raise ValueError("保存后数据与内存逐单元格不一致：" + name + repr(mismatch))
            if any(c.data_type == "f" for row in ws for c in row):
                raise ValueError("外部文本不得形成公式")
            if ws["A3"].value != snapshot["status"]:
                raise ValueError("进度状态不符")
            for r, row in enumerate(expected, 7):
                for c, value in enumerate(row, 1):
                    if isinstance(value, str) and value.startswith("https://"):
                        if not ws.cell(r, c).hyperlink or ws.cell(r, c).hyperlink.target != value:
                            raise ValueError("链接未正确保存")
        for row in wb["商品汇总"].iter_rows(min_row=7):
            if any(row[i].data_type != "s" or row[i].number_format != "@" for i in (3,4)):
                raise ValueError("ID 必须保存为文本")
        for name, columns in (("SKU明细", (2,3,5)), ("图片明细", (2,3))):
            for row in wb[name].iter_rows(min_row=7):
                if row[0].value is not None and any(row[i].data_type != "s" or row[i].number_format != "@" for i in columns):
                    raise ValueError("明细 ID 必须保存为文本")
        if wb["抓取核验"]["B23"].value != "未完成":
            raise ValueError("进度表不能标记全量通过")
    finally:
        wb.close()
    return snapshot["counts"]

def write_snapshot(payload, project_root, worktree, *, renderer=None):
    snapshot = build_snapshot(payload)
    root, daily = _safe_path(Path(project_root)), _safe_path(Path(worktree))
    target = _target_workbook(root, daily).with_name(PROGRESS_NAME)
    started = _timestamp(payload["started_at"], "started_at")
    if daily.name != started.astimezone(ZoneInfo("Asia/Shanghai")).strftime("%Y%m%d_ugreen_topsales"):
        raise ValueError("进度表目录不匹配本轮启动日")
    _safe_path(target)
    with _exclusive_publish(root):
        _check_destination(target.with_name("ugreen_topsales.xlsx"))
        _check_existing(target, snapshot)
        target.parent.mkdir(exist_ok=True)
        descriptor, name = tempfile.mkstemp(prefix=".ugreen-progress-", suffix=".xlsx", dir=target.parent)
        os.close(descriptor)
        candidate = Path(name)
        inspect_sidecar = candidate.with_name(candidate.name + ".inspect.ndjson")
        try:
            script = Path(renderer) if renderer else Path(__file__).with_name("checkpoint_excel.mjs")
            proc = subprocess.run([shutil.which("node") or "node", str(script), str(candidate)],
                                  input=json.dumps(snapshot, ensure_ascii=False, allow_nan=False), text=True,
                                  capture_output=True, timeout=120)
            if proc.returncode:
                raise RuntimeError("进度表写入失败：" + proc.stderr[-1000:])
            # artifact-tool may emit an inspection sidecar beside the requested
            # workbook. It contains the captured rows and is never an allowed
            # artifact, so remove it before validating or publishing the Excel.
            inspect_sidecar.unlink(missing_ok=True)
            add_hyperlinks(candidate, snapshot)
            validate_snapshot(candidate, snapshot)
            with candidate.open("rb") as handle:
                os.fsync(handle.fileno())
            os.replace(candidate, target)
        finally:
            candidate.unlink(missing_ok=True)
            inspect_sidecar.unlink(missing_ok=True)
    return dict(path=str(target), sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
                collection_mode=MODE, audit_status="未完成", **snapshot["counts"])


def _check_existing(target, snapshot):
    _safe_path(target)
    if target.exists():
        existing = load_workbook(target, read_only=True)
        try:
            if (tuple(existing.sheetnames) != excel.SHEET_NAMES
                    or existing["抓取核验"]["A2"].value != "UGREEN Top Sales 抓取核验（采集进度）"
                    or existing["抓取核验"]["A4"].value != "启动时间：" + snapshot["started_at"]):
                raise ValueError("已存在其他轮次的进度表，不覆盖")
            old_done = existing["抓取核验"]["B15"].value
            if not isinstance(old_done, int) or old_done > snapshot["counts"]["detail_count"]:
                raise ValueError("拒绝覆盖更多已保存详情的进度表")
            old_rows = [list(row) for row in existing["商品汇总"].iter_rows(min_row=7, values_only=True)
                        if any(value is not None for value in row)]
            new_rows = snapshot["sheets"][0]
            old_ids = [row[4] for row in old_rows]
            new_ids = [row[4] for row in new_rows]
            if old_ids != new_ids[:len(old_ids)]:
                raise ValueError("拒绝丢失、重排或替换已保存的列表身份")
            immutable_list_columns = (0, 1, 2, 3, 4, 6, 7, 8, 9, 13)
            completed_ids = set()
            for old, new in zip(old_rows, new_rows):
                if old[14] == "已完成详情":
                    completed_ids.add(old[4])
                    if old != new:
                        raise ValueError("拒绝改写已保存的完整详情汇总")
                elif old[14] == "未采集详情（标题来自列表）":
                    if any(old[index] != new[index] for index in immutable_list_columns):
                        raise ValueError("拒绝改写已保存的列表事实")
                else:
                    raise ValueError("已有进度表的详情状态无效")
            for sheet_index, name in ((1, "SKU明细"), (2, "图片明细")):
                old_detail_rows = [list(row) for row in existing[name].iter_rows(min_row=7, values_only=True)
                                   if any(value is not None for value in row) and row[3] in completed_ids]
                new_detail_rows = [list(row) for row in snapshot["sheets"][sheet_index]
                                   if row[3] in completed_ids]
                if old_detail_rows != new_detail_rows:
                    raise ValueError("拒绝改写已保存详情的 SKU 或图库")
        finally:
            existing.close()


class CheckpointSaver:
    """Save cumulative data at 50/100/... completed details and on stop/final tail."""

    def __init__(self, snapshot_factory, save, *, saved_count=0):
        if type(saved_count) is not int or saved_count < 0:
            raise ValueError("saved_count 必须是已成功保存的详情数")
        self.snapshot_factory, self.save = snapshot_factory, save
        self.saved_count = saved_count

    def maybe_save(self, *, force=False):
        payload = self.snapshot_factory()
        count = len(payload["details"])
        if count < self.saved_count:
            raise ValueError("内存详情数少于已保存数，停止以防数据丢失")
        if not force and count // 50 <= self.saved_count // 50:
            return None
        receipt = self.save(payload)
        if receipt["detail_count"] != count:
            raise ValueError("Excel 保存回执与内存详情数不一致")
        self.saved_count = count
        return receipt
