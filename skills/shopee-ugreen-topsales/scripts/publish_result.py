#!/usr/bin/env python3
"""将现有浏览器工具的完整现场数据离线校验、转换并原子发布；不访问浏览器或网络。

完整性声明必须由上游当前页面证据支持。本模块验证数据内部契约，不能仅凭 JSON
证明网站的 model/图库已经全部读取。临时候选工作簿只在发布期间存在，失败即清理。
"""

from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import datetime
from decimal import Decimal
import fcntl
import hashlib
import json
import math
import os
from pathlib import Path
import re
import stat
import sys
import tempfile
from urllib.parse import parse_qs, unquote, urlsplit
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

import excel
from prepare_daily_worktree import WorktreeError, _git
from validate_result import _target_workbook


SCHEMA_VERSION = "1.0"
STORE_URL = "https://shopee.ph/ugreen.ph?page=0&sortBy=sales&tab=0"
SHOP_ID = "64922227"
SCRIPT_RELATIVE = Path("skills/shopee-ugreen-topsales/scripts/publish_result.py")
TEMPLATE_RELATIVE = Path("skills/shopee-ugreen-topsales/assets/ugreen_topsales.xlsx")
SHANGHAI = ZoneInfo("Asia/Shanghai")
_HEADERS = (excel.SUMMARY_HEADERS, excel.SKU_HEADERS, excel.IMAGE_HEADERS, excel.AUDIT_HEADERS)
_MONTHLY = re.compile(r"(?P<n>(?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(?:\.[0-9]+)?)(?P<unit>[KM]?)(?P<plus>\+?) Sold/Month", re.I)
_IMAGE_HOSTS = {"down-ph.img.susercontent.com", "cf.shopee.ph"}


def _object(value, fields: set[str], label: str) -> dict:
    if type(value) is not dict or set(value) != fields:
        raise ValueError(f"{label}字段必须恰为：{', '.join(sorted(fields))}")
    return value


def _text(value, label: str) -> str:
    if (not isinstance(value, str) or not value.strip() or len(value) > 32767
            or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)):
        raise ValueError(f"{label}必须为非空有效文本")
    return value


def _integer(value, label: str, minimum: int = 0) -> int:
    if type(value) is not int or value < minimum:
        raise ValueError(f"{label}必须为不小于 {minimum} 的整数")
    return value


def _https(value, label: str):
    text = _text(value, label)
    if text != text.strip() or re.search(r"\s|\\", text):
        raise ValueError(f"{label} URL 含空白或反斜杠")
    try:
        parsed = urlsplit(text)
        valid = (parsed.scheme == "https" and parsed.hostname and not parsed.username
                 and not parsed.password and parsed.port is None and not parsed.fragment
                 and "@" not in parsed.netloc)
    except ValueError as exc:
        raise ValueError(f"{label} URL 无效") from exc
    if not valid:
        raise ValueError(f"{label}必须为无凭据、无端口的 HTTPS URL")
    return parsed


def _list_url(value, page_index: int) -> str:
    parsed = _https(value, "列表")
    query = parse_qs(parsed.query, keep_blank_values=True)
    if not (parsed.hostname == "shopee.ph" and parsed.path.rstrip("/") == "/ugreen.ph"
            and query.get("page") == [str(page_index)] and query.get("sortBy") == ["sales"]
            and query.get("tab") == ["0"] and query.get("shop", [SHOP_ID]) == [SHOP_ID]):
        raise ValueError("列表 URL 必须匹配 PH UGREEN、Top Sales 和当前页码")
    return value


def _identity(value) -> tuple[str, str]:
    parsed = _https(value, "商品")
    if parsed.hostname != "shopee.ph":
        raise ValueError("商品 URL 必须属于 shopee.ph")
    path = unquote(parsed.path)
    match = re.fullmatch(r"/product/([0-9]+)/([0-9]+)/?", path)
    if match is None:
        match = re.search(r"(?:-i\.|/i\.)([0-9]+)\.([0-9]+)/?$", path)
    if match is None or match[1] != SHOP_ID or int(match[2]) < 1:
        raise ValueError("商品 URL 必须绑定 PH UGREEN 店铺 64922227 的有效商品")
    return SHOP_ID, str(int(match[2]))


def _image_url(value) -> str:
    parsed = _https(value, "商品图片")
    if (parsed.hostname not in _IMAGE_HOSTS or parsed.query
            or re.fullmatch(r"/file/[A-Za-z0-9][A-Za-z0-9_.-]*", parsed.path) is None
            or parsed.path.endswith(("_tn", "_t"))):
        raise ValueError("图片必须是支持的 PH 商品 CDN /file/ 完整 URL，不能是缩略路径或非商品页面")
    return value


def _timestamp(value, label: str) -> datetime:
    _text(value, label)
    try:
        timestamp = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"{label}必须是带时区 ISO 时间") from exc
    if timestamp.tzinfo is None or timestamp.utcoffset() is None:
        raise ValueError(f"{label}必须是带时区 ISO 时间")
    return timestamp


def _monthly(value) -> tuple[str | None, int | None, str]:
    if value is None:
        return None, None, "not_displayed"
    if not isinstance(value, str) or (match := _MONTHLY.fullmatch(value)) is None:
        raise ValueError("月销必须为完整的 'N Sold/Month' 原文或 null")
    unit = match["unit"].upper()
    if "." in match["n"] and not unit:
        raise ValueError("无 K/M 单位的月销必须是整数")
    amount = Decimal(match["n"].replace(",", "")) * {"": 1, "K": 1000, "M": 1000000}[unit]
    if amount != amount.to_integral_value() or amount > 2**53 - 1:
        raise ValueError("月销下界必须是 Excel 可精确保存的整数")
    return value, int(amount), "displayed"


def _convert(payload: dict) -> tuple[dict, datetime]:
    """完全从现场载荷重算审计；不接受外来排名、计数或 audit。"""
    _object(payload, {"schema_version", "store_url", "started_at", "captured_at", "pages", "details"}, "载荷")
    if payload["schema_version"] != SCHEMA_VERSION:
        raise ValueError("只支持 schema_version='1.0'")
    _list_url(payload["store_url"], 0)
    started = _timestamp(payload["started_at"], "started_at")
    captured = _timestamp(payload["captured_at"], "captured_at")
    if captured < started:
        raise ValueError("captured_at 不能早于 started_at")
    pages, details = payload["pages"], payload["details"]
    if type(pages) is not list or not pages or type(details) is not list:
        raise ValueError("pages 必须为非空数组，details 必须为数组")
    cards_by_id, occurrences = {}, 0
    for index, page in enumerate(pages):
        _object(page, {"page_index", "total_pages", "url", "next_enabled", "cards"}, "列表页")
        if (_integer(page["page_index"], "page_index") != index
                or _integer(page["total_pages"], "total_pages", 1) != len(pages)):
            raise ValueError("列表页必须从 0 连续排列，total_pages 必须一致且已全部读取")
        _list_url(page["url"], index)
        if type(page["next_enabled"]) is not bool or page["next_enabled"] != (index < len(pages) - 1):
            raise ValueError("列表 next_enabled 与末页状态不一致")
        cards = page["cards"]
        if type(cards) is not list or not 1 <= len(cards) <= 30 or (index < len(pages) - 1 and len(cards) != 30):
            raise ValueError("非末页必须恰有 30 张商品卡，末页必须有 1–30 张")
        page_ids, new_count = set(), 0
        for position, card in enumerate(cards, 1):
            _object(card, {"product_url", "title", "price_php", "monthly_sales_text"}, "商品卡")
            key = _identity(card["product_url"])
            if key in page_ids:
                raise ValueError("单页商品卡存在重复身份")
            page_ids.add(key)
            _text(card["title"], "列表标题")
            price = card["price_php"]
            if (type(price) not in (int, float) or price < 0 or price > 2**53 - 1
                    or not math.isfinite(price)):
                raise ValueError("price_php 必须为非负有限数值")
            monthly, lower, observation = _monthly(card["monthly_sales_text"])
            occurrences += 1
            if key not in cards_by_id:
                new_count += 1
                cards_by_id[key] = {
                    "global_rank": len(cards_by_id) + 1, "source_page": index,
                    "source_position": position, "shop_id": key[0], "item_id": key[1],
                    "price_php": price, "monthly_sales_text": monthly,
                    "monthly_sales_count_lower_bound": lower, "monthly_sales_observation": observation,
                }
        if new_count == 0:
            raise ValueError("列表页没有新身份，不能用重复页证明分页完整")
    details_by_id = {}
    for detail in details:
        _object(detail, {"product_url", "title", "skus", "main_image_url", "secondary_image_urls", "models_complete", "gallery_complete"}, "商品详情")
        key = _identity(detail["product_url"])
        if key in details_by_id or key not in cards_by_id:
            raise ValueError("详情包含重复或列表范围外商品")
        if detail["models_complete"] is not True or detail["gallery_complete"] is not True:
            raise ValueError("完整 model 集合和图库均必须有本轮证据确认")
        _text(detail["title"], "详情标题")
        if type(detail["skus"]) is not list or not detail["skus"]:
            raise ValueError("每个商品必须有完整非空 SKU/model 数组")
        models, skus = set(), []
        for sku in detail["skus"]:
            _object(sku, {"model_id", "sku_props", "sku_image_url"}, "SKU")
            model = sku["model_id"]
            if not isinstance(model, str) or re.fullmatch(r"[0-9]+", model) is None or int(model) in models:
                raise ValueError("model_id 必须为唯一的纯数字文本")
            models.add(int(model))
            props = sku["sku_props"]
            if type(props) is not dict or not props:
                raise ValueError("sku_props 必须是非空规格映射，不能猜造组合")
            for name, option in props.items():
                _text(name, "规格名称")
                _text(option, "规格选项")
            image_url = sku["sku_image_url"]
            if image_url is not None:
                _image_url(image_url)
            skus.append({"model_id": model, "sku_props": dict(props), "sku_image_url": image_url})
        main_image = _image_url(detail["main_image_url"])
        if type(detail["secondary_image_urls"]) is not list:
            raise ValueError("secondary_image_urls 必须为数组")
        secondary = [_image_url(value) for value in detail["secondary_image_urls"]]
        if len({main_image, *secondary}) != 1 + len(secondary):
            raise ValueError("主图/副图库存在重复 URL")
        details_by_id[key] = {"product_url": detail["product_url"], "title": detail["title"],
                              "skus": skus, "main_image_url": main_image, "secondary_image_urls": secondary}
    if set(details_by_id) != set(cards_by_id):
        raise ValueError("详情缺失，必须逐个覆盖本轮列表全部唯一商品")
    products = [{**card, **details_by_id[key]} for key, card in cards_by_id.items()]
    audit = {"list_page_count": len(pages), "list_input_product_occurrences": occurrences,
             "list_duplicate_occurrence_count": occurrences - len(products),
             "listed_product_count": len(products), "detail_success_count": len(products),
             "detail_failure_count": 0, "detail_parse_success_count": len(products),
             "detail_parse_error_count": 0,
             "products_without_monthly_sales_display": sum(p["monthly_sales_observation"] == "not_displayed" for p in products)}
    return {"store_url": STORE_URL, "captured_at": captured.astimezone(SHANGHAI).strftime("%Y/%m/%d %H:%M:%S"),
            "products": products, "audit": audit, "collection_mode": "full_topsales"}, started


def _safe_path(path: Path) -> Path:
    path = Path(path).expanduser()
    if ".." in path.parts:
        raise ValueError("路径不能包含 ..")
    path = path.absolute()
    for part in (path, *path.parents):
        if part.is_symlink():
            raise ValueError(f"路径不能含符号链接：{part}")
    return path


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_template_contract(project_root: Path) -> dict:
    """只读规范表名/表头及文件摘要，不以旧商品业务数据填充本轮结果。"""
    template = _safe_path(Path(project_root) / TEMPLATE_RELATIVE)
    if not template.is_file():
        raise ValueError(f"缺少只读 Excel 规范：{template}")
    digest = _sha256(template)
    try:
        workbook = load_workbook(template, read_only=True, data_only=False, keep_links=False)
        try:
            if tuple(workbook.sheetnames) != excel.SHEET_NAMES:
                raise ValueError("规范工作簿四表名称或顺序与导出 schema 不一致")
            for name, expected in zip(excel.SHEET_NAMES, _HEADERS):
                sheet = workbook[name]
                row = next(sheet.iter_rows(min_row=excel.HEADER_ROW, max_row=excel.HEADER_ROW, values_only=True))
                if tuple(row) != expected:
                    raise ValueError(f"规范表 {name} 的表头与导出 schema 不一致")
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取有效的只读 Excel 规范") from exc
    if _sha256(template) != digest:
        raise ValueError("读取期间 Excel 规范发生变化")
    return {"path": str(template), "sha256": digest, "sheets": list(excel.SHEET_NAMES), "header_row": excel.HEADER_ROW}


def _check_destination(target: Path, candidate: Path | None = None) -> None:
    _safe_path(target)
    if target.exists() and not target.is_file():
        raise ValueError("正式结果路径不是普通文件")
    if target.parent.exists():
        if not target.parent.is_dir():
            raise ValueError("result 路径不是目录")
        allowed = {target} | ({candidate} if candidate is not None else set())
        progress = _safe_path(target.with_name("ugreen_topsales_progress.xlsx"))
        if progress.exists():
            if not progress.is_file():
                raise ValueError("进度 Excel 路径不是普通文件")
            allowed.add(progress)
        if any(entry not in allowed for entry in target.parent.iterdir()):
            raise ValueError("result/ 含未知文件；不会自动删除或忽略")
    if candidate is not None:
        _safe_path(candidate)
        if not candidate.is_file() or candidate.parent != target.parent:
            raise ValueError("本次临时候选路径异常")


def _gates(root: Path, daily: Path, started: datetime, candidate: Path | None = None) -> tuple[Path, str, str]:
    _safe_path(root)
    _safe_path(daily)
    _safe_path(root / ".git")
    target = _target_workbook(root, daily)
    compact = started.astimezone(SHANGHAI).strftime("%Y%m%d")
    if daily.name != f"{compact}_ugreen_topsales":
        raise ValueError("日期 worktree 必须匹配 started_at 的上海启动日")
    script = _safe_path(Path(__file__))
    if script.resolve() != (daily / SCRIPT_RELATIVE).resolve():
        raise ValueError("发布代码必须从目标日期 worktree 执行")
    branch = _git(daily, "symbolic-ref", "--quiet", "--short", "HEAD").stdout.strip()
    if branch != f"daily/{compact}-ugreen-topsales":
        raise ValueError("日期 worktree 分支与启动日不匹配")
    for label, path in (("主项目", root), ("日期 worktree", daily)):
        if _git(path, "status", "--porcelain", "--untracked-files=all").stdout.strip():
            raise ValueError(f"{label}存在未提交改动；不会自动处理 Git")
        _git(path, "ls-files", "--error-unmatch", "--", str(TEMPLATE_RELATIVE))
    head = _git(root, "rev-parse", "HEAD").stdout.strip()
    if head != _git(daily, "rev-parse", "HEAD").stdout.strip():
        raise ValueError("主项目和日期 worktree 必须为同一已提交 HEAD")
    root_spec = validate_template_contract(root)
    daily_spec = validate_template_contract(daily)
    if root_spec["sha256"] != daily_spec["sha256"]:
        raise ValueError("主项目和日期 worktree 的只读规范 SHA-256 不一致")
    _check_destination(target, candidate)
    return target, root_spec["sha256"], head


@contextmanager
def _exclusive_publish(root: Path):
    lock_path = _safe_path(root / ".git/ugreen_daily_run.lock")
    descriptor = os.open(lock_path, os.O_CREAT | os.O_RDWR | os.O_NOFOLLOW, 0o600)
    try:
        if not stat.S_ISREG(os.fstat(descriptor).st_mode):
            raise ValueError("任务锁必须为普通文件")
        try:
            fcntl.flock(descriptor, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise ValueError("已有另一个全量采集或发布任务正在运行") from exc
        yield
    finally:
        os.close(descriptor)


def publish(payload: dict, project_root: Path, worktree: Path) -> dict:
    """严格校验现场载荷，在 clean 同 HEAD 的启动日 worktree 原子发布。"""
    result, started = _convert(payload)
    root, daily = _safe_path(project_root), _safe_path(worktree)
    # 注册关系与执行位置先于创建锁/结果目录检查，拒绝越界路径不产生文件。
    target, template_hash, head = _gates(root, daily, started)
    with _exclusive_publish(root):
        if _gates(root, daily, started) != (target, template_hash, head):
            raise ValueError("取得任务锁前发布前置状态发生变化")
        target.parent.mkdir(exist_ok=True)
        descriptor, name = tempfile.mkstemp(prefix=".ugreen-publish-", suffix=".xlsx", dir=target.parent)
        os.close(descriptor)
        candidate = Path(name)
        try:
            excel.write_excel(result, candidate)
            stats = excel.validate_excel(candidate)
            if stats.audit_status != "通过" or stats.collection_mode != "full_topsales":
                raise ValueError("独立工作簿审计未通过全量 Top Sales 契约")
            candidate_hash = _sha256(candidate)
            if _gates(root, daily, started, candidate) != (target, template_hash, head):
                raise ValueError("发布前 Git HEAD 或只读 Excel 规范发生变化")
            os.replace(candidate, target)
            # write_excel 已 fsync 候选内容；这里只同步最终目录项。
            try:
                directory_fd = os.open(target.parent, os.O_RDONLY)
                try:
                    os.fsync(directory_fd)
                finally:
                    os.close(directory_fd)
            except OSError:
                pass  # 部分文件系统不支持目录 fsync，原子替换已完成。
        finally:
            if candidate.exists():
                candidate.unlink()
    return {"path": str(target), "sha256": candidate_hash, "template_sha256": template_hash,
            "schema_version": SCHEMA_VERSION, "collection_mode": stats.collection_mode,
            "list_page_count": result["audit"]["list_page_count"], "product_count": stats.product_count,
            "sku_count": stats.sku_count, "main_image_count": stats.main_image_count,
            "secondary_image_count": stats.secondary_image_count, "audit_status": stats.audit_status}


def _json_object(pairs):
    value = {}
    for key, item in pairs:
        if key in value:
            raise ValueError(f"JSON 对象包含重复键：{key}")
        value[key] = item
    return value


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--worktree", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        payload = json.load(sys.stdin, object_pairs_hook=_json_object)
        receipt = publish(payload, args.project_root, args.worktree)
    except (ValueError, OSError, WorktreeError) as exc:
        print(f"发布失败：{exc}", file=sys.stderr)
        return 2
    print(json.dumps(receipt, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
